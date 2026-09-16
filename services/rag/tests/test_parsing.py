"""Parser tests, run against real generated files.

Regenerate the fixtures with: python tests/fixtures/generate.py
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.models import BlockType, DocumentKind
from app.parsing.pdf import PyMuPDFParser
from app.parsing.pptx import PythonPptxParser

FIXTURES = Path(__file__).parent / "fixtures"

pytestmark = pytest.mark.skipif(
    not (FIXTURES / "normalization.pdf").exists(),
    reason="fixtures not generated; run tests/fixtures/generate.py",
)


@pytest.fixture
async def pdf_doc():
    return await PyMuPDFParser().parse(
        FIXTURES / "normalization.pdf",
        document_id="doc-pdf",
        filename="normalization.pdf",
    )


@pytest.fixture
async def pptx_doc():
    return await PythonPptxParser().parse(
        FIXTURES / "indexing.pptx",
        document_id="doc-pptx",
        filename="indexing.pptx",
    )


class TestPdfParser:
    def test_declares_the_formats_it_handles(self) -> None:
        parser = PyMuPDFParser()
        assert parser.supports(DocumentKind.PDF)
        assert not parser.supports(DocumentKind.PPTX)

    async def test_reports_page_count(self, pdf_doc) -> None:
        assert pdf_doc.page_count == 2
        assert pdf_doc.parser == "pymupdf"

    async def test_recovers_the_heading_hierarchy(self, pdf_doc) -> None:
        headings = [b for b in pdf_doc.blocks if b.type == BlockType.HEADING]
        texts = [b.text for b in headings]

        assert "Database Normalization" in texts
        assert "Third Normal Form" in texts

        # The document title is set larger than the section headings, so it
        # must come out one level shallower.
        title = next(b for b in headings if b.text == "Database Normalization")
        section = next(b for b in headings if b.text == "Third Normal Form")
        assert title.level is not None
        assert section.level is not None
        assert title.level < section.level

    async def test_assigns_correct_page_numbers(self, pdf_doc) -> None:
        # Citations are only useful if the page is right.
        by_text = {b.text[:30]: b for b in pdf_doc.blocks}
        first_nf = next(
            b for k, b in by_text.items() if k.startswith("First Normal Form")
        )
        third_nf = next(
            b for k, b in by_text.items() if k.startswith("Third Normal Form")
        )
        assert first_nf.page_number == 1
        assert third_nf.page_number == 2

    async def test_separates_paragraphs_by_vertical_spacing(self, pdf_doc) -> None:
        # A PDF records no paragraph markers. Two consecutive paragraphs under
        # one heading must still come out as two blocks.
        paragraphs = [b for b in pdf_doc.blocks if b.type == BlockType.PARAGRAPH]
        third_nf = [p for p in paragraphs if "third normal form" in p.text.lower()]
        bcnf = [p for p in paragraphs if "Boyce-Codd" in p.text]

        assert len(third_nf) == 1
        assert len(bcnf) == 1
        assert third_nf[0] is not bcnf[0]

    async def test_extracts_the_table(self, pdf_doc) -> None:
        tables = [b for b in pdf_doc.blocks if b.type == BlockType.TABLE]
        assert len(tables) == 1
        assert "Repeating groups" in tables[0].text

    async def test_table_text_is_not_also_emitted_as_prose(self, pdf_doc) -> None:
        # Table regions are masked out of the prose stream. Without that,
        # every cell is indexed twice and pollutes retrieval.
        #
        # "Eliminates" is a column header that appears nowhere in the body
        # text, so it isolates duplication from legitimate overlap -- phrases
        # like "Repeating groups" occur in both the table and the prose.
        prose = " ".join(
            b.text for b in pdf_doc.blocks if b.type != BlockType.TABLE
        )
        assert "Eliminates" not in prose
        assert "Transitive dependency | 2NF" not in prose

    async def test_blocks_are_in_reading_order(self, pdf_doc) -> None:
        orders = [b.order for b in pdf_doc.blocks]
        assert orders == sorted(orders)

        # The table sits at the foot of page 2, so it must come last -- not
        # first, which is the order tables are detected in.
        assert pdf_doc.blocks[-1].type == BlockType.TABLE

    async def test_a_pdf_with_no_text_yields_no_blocks(self) -> None:
        # The scanned-document case. It must return cleanly so the pipeline can
        # report "no extractable text" rather than crash.
        doc = await PyMuPDFParser().parse(
            FIXTURES / "empty.pdf", document_id="d", filename="empty.pdf"
        )
        assert doc.page_count == 2
        assert doc.blocks == []


class TestPptxParser:
    def test_declares_the_formats_it_handles(self) -> None:
        parser = PythonPptxParser()
        assert parser.supports(DocumentKind.PPTX)
        assert parser.supports(DocumentKind.PPT)
        assert not parser.supports(DocumentKind.PDF)

    async def test_reports_slide_count(self, pptx_doc) -> None:
        assert pptx_doc.page_count == 4
        assert pptx_doc.metadata["slide_count"] == 4

    async def test_every_block_carries_a_slide_number(self, pptx_doc) -> None:
        assert all(b.slide_number is not None for b in pptx_doc.blocks)
        assert all(b.page_number is None for b in pptx_doc.blocks)

    async def test_slide_titles_become_top_level_headings(self, pptx_doc) -> None:
        headings = [b for b in pptx_doc.blocks if b.type == BlockType.HEADING]
        assert [h.text for h in headings] == [
            "Introduction to Indexing",
            "B-Tree Indexes",
            "Composite Indexes",
            "Index Comparison",
        ]
        assert all(h.level == 1 for h in headings)

    async def test_extracts_speaker_notes(self, pptx_doc) -> None:
        # Lecture slides are terse; the sentence that actually explains the
        # bullet often lives only in the notes.
        notes = [b for b in pptx_doc.blocks if b.type == BlockType.SPEAKER_NOTE]
        assert len(notes) == 2
        assert any("buying read speed" in n.text for n in notes)

    async def test_preserves_list_nesting(self, pptx_doc) -> None:
        items = [b for b in pptx_doc.blocks if b.type == BlockType.LIST_ITEM]
        assert items
        assert all(b.level is not None and b.level > 1 for b in items)

    async def test_extracts_slide_tables(self, pptx_doc) -> None:
        tables = [b for b in pptx_doc.blocks if b.type == BlockType.TABLE]
        assert len(tables) == 1
        assert "B-Tree | Yes | Yes" in tables[0].text

    async def test_notes_follow_the_slide_they_belong_to(self, pptx_doc) -> None:
        for index, block in enumerate(pptx_doc.blocks):
            if block.type == BlockType.SPEAKER_NOTE:
                previous = pptx_doc.blocks[index - 1]
                assert previous.slide_number == block.slide_number


def _png(path: Path, size: tuple[int, int] = (320, 240)) -> Path:
    from PIL import Image

    Image.new("RGB", size, (40, 90, 160)).save(path)
    return path


def _deck(tmp_path: Path, build) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    build(presentation, _png(tmp_path / "diagram.png"), Inches)
    out = tmp_path / "deck.pptx"
    presentation.save(str(out))
    return out


async def _parse(path: Path):
    return await PythonPptxParser().parse(
        path, document_id="doc-img", filename=path.name
    )


class TestPptxImages:
    """Every picture reaches the vision stage, however it was placed."""

    async def test_a_free_standing_picture_is_extracted(self, tmp_path) -> None:
        def build(p, png, inches):
            slide = p.slides.add_slide(p.slide_layouts[5])
            slide.shapes.title.text = "Architecture"
            slide.shapes.add_picture(str(png), inches(1), inches(2))

        doc = await _parse(_deck(tmp_path, build))

        assert [(i.slide_number, i.width, i.height) for i in doc.images] == [
            (1, 320, 240)
        ]

    async def test_a_picture_in_a_placeholder_is_extracted(self, tmp_path) -> None:
        # Inserting through a placeholder is PowerPoint's default, and the
        # shape then reports as PLACEHOLDER rather than PICTURE.
        def build(p, png, inches):
            slide = p.slides.add_slide(p.slide_layouts[8])
            slide.shapes.title.text = "Proposed block diagram"
            picture = slide.placeholders[1].insert_picture(str(png))
            assert picture.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER

        doc = await _parse(_deck(tmp_path, build))

        assert len(doc.images) == 1
        assert doc.images[0].slide_number == 1
        assert doc.images[0].context_hint.startswith("Proposed block diagram")

    async def test_pictures_and_text_inside_a_group_are_read(self, tmp_path) -> None:
        def build(p, png, inches):
            slide = p.slides.add_slide(p.slide_layouts[5])
            slide.shapes.title.text = "Signal path"
            group = slide.shapes.add_group_shape()
            group.shapes.add_picture(str(png), inches(1), inches(2))
            label = group.shapes.add_textbox(inches(5), inches(2), inches(3), inches(1))
            label.text_frame.text = "Microphone converts vibration to voltage"

        doc = await _parse(_deck(tmp_path, build))

        assert len(doc.images) == 1
        assert any(
            "Microphone converts vibration" in b.text
            for b in doc.blocks
            if b.type != BlockType.HEADING
        )

    async def test_a_slide_without_pictures_yields_no_images(self, pptx_doc) -> None:
        assert pptx_doc.images == []
