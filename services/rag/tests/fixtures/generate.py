"""Generates the test fixtures.

Real files, not mocks. Parsing bugs live in the gap between what a library is
documented to return and what it actually returns for a given file, and a mock
reproduces the documentation rather than the behaviour.

Run: python tests/fixtures/generate.py
"""

from __future__ import annotations

from pathlib import Path
from typing import NamedTuple

import fitz
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent


class Item(NamedTuple):
    """One piece of fixture content.

    Fields are named because an earlier positional version conflated heading
    level with page number, which silently placed content on the wrong page.
    """

    kind: str  # "heading" | "body"
    text: str
    page: int
    level: int | None = None


# Deliberately shaped like real study material: a heading hierarchy, technical
# terms and abbreviations that only keyword search will match exactly, a
# definition worth citing, and a table.
PDF_CONTENT: list[Item] = [
    Item("heading", "Database Normalization", page=1, level=1),
    Item(
        "body",
        "Normalization is the process of organizing data in a database to "
        "reduce redundancy and improve data integrity. It was introduced by "
        "Edgar F. Codd in 1970 as part of the relational model.",
        page=1,
    ),
    Item("heading", "First Normal Form", page=1, level=2),
    Item(
        "body",
        "A relation is in first normal form (1NF) if every attribute contains "
        "only atomic values. Repeating groups are not permitted. This is the "
        "minimum requirement for a relational table.",
        page=1,
    ),
    Item("heading", "Second Normal Form", page=1, level=2),
    Item(
        "body",
        "A relation is in second normal form (2NF) if it is in 1NF and every "
        "non-prime attribute is fully functionally dependent on the whole of "
        "every candidate key. Partial dependencies are eliminated at this "
        "stage.",
        page=1,
    ),
    Item("heading", "Third Normal Form", page=2, level=2),
    Item(
        "body",
        "A relation is in third normal form (3NF) if it is in 2NF and no "
        "non-prime attribute is transitively dependent on any candidate key. "
        "Transitive dependency means an attribute depends on another "
        "non-prime attribute rather than directly on the key.",
        page=2,
    ),
    Item(
        "body",
        "The Boyce-Codd normal form (BCNF) is a stricter version of 3NF. "
        "Every determinant must be a candidate key. Most relations in 3NF are "
        "also in BCNF, but exceptions exist when a relation has multiple "
        "overlapping candidate keys.",
        page=2,
    ),
    Item("heading", "Denormalization", page=2, level=2),
    Item(
        "body",
        "Denormalization deliberately introduces redundancy to improve read "
        "performance. It trades write cost and integrity risk for query speed, "
        "and is applied only after measurement shows a real bottleneck.",
        page=2,
    ),
]

PPTX_SLIDES: list[tuple[str, list[str], str]] = [
    (
        "Introduction to Indexing",
        [
            "An index is a data structure that improves lookup speed",
            "Indexes cost storage and slow down writes",
            "The query planner decides whether to use one",
        ],
        "An index is a trade. You are buying read speed with write speed and "
        "disk. The exam question is almost always about when that trade is "
        "worth making, not about how a B-tree works internally.",
    ),
    (
        "B-Tree Indexes",
        [
            "Balanced tree structure with logarithmic lookup",
            "Supports equality and range queries",
            "The default index type in PostgreSQL",
        ],
        "B-tree is the default for a reason: it handles both equality and "
        "range predicates. Hash indexes only handle equality, which is why "
        "they are rarely worth using.",
    ),
    (
        "Composite Indexes",
        [
            "An index on multiple columns in a defined order",
            "Column order determines which queries can use it",
            "Leftmost prefix rule applies",
        ],
        "",
    ),
]


def _wrap(text: str, chars_per_line: int) -> list[str]:
    """Greedy word wrap. Deterministic, so fixtures are byte-stable."""
    lines: list[str] = []
    current: list[str] = []
    length = 0
    for word in text.split():
        if current and length + 1 + len(word) > chars_per_line:
            lines.append(" ".join(current))
            current, length = [word], len(word)
        else:
            current.append(word)
            length += (1 if length else 0) + len(word)
    if current:
        lines.append(" ".join(current))
    return lines


def build_pdf(path: Path) -> None:
    """A PDF with a real heading hierarchy expressed through font size.

    Text is placed line by line rather than through insert_textbox, so nothing
    silently overflows off the page -- a fixture that quietly loses content
    makes every assertion written against it meaningless.
    """
    doc = fitz.open()
    sizes = {1: 20.0, 2: 15.0}
    body_size = 11.0
    margin = 72.0
    bottom_limit = 560.0  # leaves room for the table at the foot of page 2

    doc.new_page()
    doc.new_page()
    cursors = {1: margin, 2: margin}

    # Pages are indexed fresh on each use rather than held in a list: PyMuPDF
    # invalidates existing Page handles when the document grows.
    for item in PDF_CONTENT:
        page = doc[item.page - 1]
        y = cursors[item.page]

        if item.kind == "heading":
            size = sizes[item.level or 2]
            fontname = "helvetica-bold"
            leading = size * 1.4
            wrapped = [item.text]
            y += 10
        else:
            size = body_size
            fontname = "helvetica"
            leading = size * 1.35
            wrapped = _wrap(item.text, 78)

        for line in wrapped:
            if y > bottom_limit:
                raise RuntimeError(
                    f"fixture overflowed page {item.page}; content would be lost"
                )
            page.insert_text(
                (margin, y), line, fontsize=size, fontname=fontname
            )
            y += leading

        cursors[item.page] = y + 6

    # A table on page 2, to exercise table extraction and de-duplication.
    page = doc[1]
    table_y = 600.0
    rows = [
        ["Normal Form", "Eliminates", "Requires"],
        ["1NF", "Repeating groups", "Atomic values"],
        ["2NF", "Partial dependency", "1NF"],
        ["3NF", "Transitive dependency", "2NF"],
    ]
    for row in rows:
        x = 72.0
        for cell in row:
            page.draw_rect(fitz.Rect(x, table_y, x + 150, table_y + 20))
            page.insert_text((x + 4, table_y + 14), cell, fontsize=9)
            x += 150
        table_y += 20

    doc.save(path)
    doc.close()


def build_pptx(path: Path) -> None:
    """A deck with titles, bullets, speaker notes and a table."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # Title and Content

    for title, bullets, notes in PPTX_SLIDES:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title

        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 1

        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # A slide carrying only a table.
    blank = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(blank)
    slide.shapes.title.text = "Index Comparison"
    shape = slide.shapes.add_table(
        3, 3, Inches(0.5), Inches(2.0), Inches(9.0), Inches(2.0)
    )
    data = [
        ["Type", "Equality", "Range"],
        ["B-Tree", "Yes", "Yes"],
        ["Hash", "Yes", "No"],
    ]
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = shape.table.cell(r, c)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    presentation.save(path)


def build_empty_pdf(path: Path) -> None:
    """A PDF with pages but no extractable text -- the scanned-document case."""
    doc = fitz.open()
    doc.new_page()
    doc.new_page()
    doc.save(path)
    doc.close()


def main() -> None:
    HERE.mkdir(parents=True, exist_ok=True)
    build_pdf(HERE / "normalization.pdf")
    build_pptx(HERE / "indexing.pptx")
    build_empty_pdf(HERE / "empty.pdf")
    for name in ("normalization.pdf", "indexing.pptx", "empty.pdf"):
        size = (HERE / name).stat().st_size
        print(f"  {name}: {size:,} bytes")


if __name__ == "__main__":
    main()
